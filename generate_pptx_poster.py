import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_a3_poster():
    prs = Presentation()
    # A3 Landscape dimensions: 420mm x 297mm -> 16.535 x 11.693 inches
    prs.slide_width = Inches(16.535)
    prs.slide_height = Inches(11.693)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # ── Color Palette ──
    BG_DARK = RGBColor(11, 19, 43)          # Deep Navy
    BG_CARD = RGBColor(255, 255, 255)       # Card White
    TEXT_DARK = RGBColor(15, 23, 42)        # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139)    # Slate 500
    TEXT_BODY = RGBColor(51, 65, 85)        # Slate 700
    PRIMARY = RGBColor(26, 86, 164)         # SBI Royal Blue
    ACCENT_GREEN = RGBColor(22, 163, 74)    # Green 600
    ACCENT_AMBER = RGBColor(217, 119, 6)    # Amber 600
    ACCENT_RED = RGBColor(220, 38, 38)      # Red 600
    BORDER_LIGHT = RGBColor(226, 232, 240)  # Slate 200
    CODE_BG = RGBColor(15, 23, 42)          # Code Slate 900
    HL_BG = RGBColor(239, 246, 255)         # Light Blue Tint
    HL_BORDER = RGBColor(37, 99, 235)       # Blue Border

    # ── Slide Background ──
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()

    # Dimensions & Margins
    margin_x = Inches(0.4)
    margin_y = Inches(0.35)
    content_width = prs.slide_width - (margin_x * 2)

    # ── Header Card ──
    header_height = Inches(1.25)
    header_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin_x, margin_y, content_width, header_height)
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = BG_CARD
    header_box.line.color.rgb = PRIMARY
    header_box.line.width = Pt(2.5)

    tf_h = header_box.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = Inches(0.25)
    tf_h.margin_top = Inches(0.12)
    tf_h.margin_right = Inches(3.8)
    tf_h.margin_bottom = Inches(0.1)

    p0 = tf_h.paragraphs[0]
    p0.text = "SECURE CI/CD PIPELINE FOR A LOAN APPLICATION TRACKING PORTAL"
    p0.font.bold = True
    p0.font.size = Pt(20)
    p0.font.color.rgb = TEXT_DARK

    p1 = tf_h.add_paragraph()
    p1.text = "Automated DevOps Architecture  •  Dual-Role (Borrower/Officer) RBAC  •  SBI Underwriting Engine  •  9-Journey Selenium Quality Gate"
    p1.font.bold = True
    p1.font.size = Pt(10.5)
    p1.font.color.rgb = PRIMARY

    # Header Meta Badges (Right side)
    meta_x = margin_x + content_width - Inches(3.6)
    meta_y = margin_y + Inches(0.12)
    meta_w = Inches(3.4)
    meta_h = Inches(1.0)
    meta_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, meta_x, meta_y, meta_w, meta_h)
    meta_box.fill.solid()
    meta_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    meta_box.line.color.rgb = RGBColor(203, 213, 225)
    meta_box.line.width = Pt(1)

    tf_m = meta_box.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = Inches(0.12)
    tf_m.margin_top = Inches(0.08)
    tf_m.margin_right = Inches(0.12)
    tf_m.margin_bottom = Inches(0.08)

    mp0 = tf_m.paragraphs[0]
    mp0.text = "LEAD ENGINEER: Vicky Vilas Pukale  |  ROLL: 23102A0033"
    mp0.font.bold = True
    mp0.font.size = Pt(8.5)
    mp0.font.color.rgb = TEXT_DARK

    mp1 = tf_m.add_paragraph()
    mp1.text = "RELEASE: v2.1.0-auth  •  AUTOMATED TESTS: 27 / 27 PASSED"
    mp1.font.bold = True
    mp1.font.size = Pt(8.5)
    mp1.font.color.rgb = PRIMARY

    mp2 = tf_m.add_paragraph()
    mp2.text = "STATUS: [✓] CI/CD PRODUCTION QUALITY GATE APPROVED"
    mp2.font.bold = True
    mp2.font.size = Pt(8.5)
    mp2.font.color.rgb = ACCENT_GREEN

    # ── 4-Column Grid Setup ──
    col_gap = Inches(0.2)
    num_cols = 4
    total_gap = col_gap * (num_cols - 1)
    col_width = (content_width - total_gap) / num_cols
    col_top = margin_y + header_height + Inches(0.2)
    footer_height = Inches(0.5)
    col_height = prs.slide_height - col_top - footer_height - Inches(0.25)

    def add_card(x, y, w, h, title, icon=""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        # Header accent bar inside card
        header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.1), y + Inches(0.08), w - Inches(0.2), Inches(0.32))
        header_bar.fill.solid()
        header_bar.fill.fore_color.rgb = RGBColor(248, 250, 252)
        header_bar.line.color.rgb = RGBColor(226, 232, 240)
        header_bar.line.width = Pt(0.75)

        tf = header_bar.text_frame
        tf.word_wrap = False
        tf.margin_left = Inches(0.08)
        tf.margin_top = Inches(0.04)
        p = tf.paragraphs[0]
        p.text = f"{icon}  {title}".strip().upper()
        p.font.bold = True
        p.font.size = Pt(9.5)
        p.font.color.rgb = PRIMARY

        # Content text frame
        content_box = slide.shapes.add_textbox(x + Inches(0.1), y + Inches(0.42), w - Inches(0.2), h - Inches(0.46))
        tf_c = content_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.04)
        tf_c.margin_right = Inches(0.04)
        tf_c.margin_top = Inches(0.04)
        tf_c.margin_bottom = Inches(0.04)
        return tf_c

    # ════════════════ COLUMN 1: DOMAIN & SBI ENGINE ════════════════
    col1_x = margin_x
    c1_h1 = Inches(2.7)
    c1_h2 = col_height - c1_h1 - Inches(0.18)

    # Card 1.1 Problem Statement
    tf1_1 = add_card(col1_x, col_top, col_width, c1_h1, "Problem Statement & Scope", "[★]")
    p = tf1_1.paragraphs[0]
    p.text = "Traditional retail loan processing is plagued by manual friction, opaque progress tracking, lack of audit separation, and unverified deployments."
    p.font.size = Pt(8.5)
    p.font.color.rgb = TEXT_BODY

    p = tf1_1.add_paragraph()
    p.text = "• Core Objective: Deliver an enterprise-grade automated lending portal featuring secure dual-role separation, automated risk calculation, and resilient CI/CD delivery."
    p.font.size = Pt(8)
    p.font.color.rgb = TEXT_BODY

    p = tf1_1.add_paragraph()
    p.text = "• Key Technologies: Spring Boot 3.2, Thymeleaf, H2 Database, Jenkins Declarative Pipeline, Docker Multi-Stage, Selenium Headless, Ansible Automation."
    p.font.size = Pt(8)
    p.font.color.rgb = TEXT_BODY

    p = tf1_1.add_paragraph()
    p.text = "• Reliability Goal: Sub-15s recovery and 100% test gate passing before deployment."
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    # Card 1.2 SBI Underwriting Engine
    tf1_2 = add_card(col1_x, col_top + c1_h1 + Inches(0.18), col_width, c1_h2, "The 4 SBI Underwriting Pillars", "[⚙]")
    p = tf1_2.paragraphs[0]
    p.text = "Live algorithmic credit risk assessment evaluating retail borrower solvency:"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = tf1_2.add_paragraph()
    p.text = "1. CIBIL Score (≥ 750 | Hard Floor 650):\n   Credit bureau rating for financial solvency."
    p.font.size = Pt(8)
    p.font.color.rgb = TEXT_BODY

    p = tf1_2.add_paragraph()
    p.text = "2. Net Monthly Income (≥ ₹20,000 / month):\n   Verifies adequate monthly debt servicing ability."
    p.font.size = Pt(8)
    p.font.color.rgb = TEXT_BODY

    p = tf1_2.add_paragraph()
    p.text = "3. Applicant Age Range (21 – 60 Years):\n   SBI standard working-age boundary."
    p.font.size = Pt(8)
    p.font.color.rgb = TEXT_BODY

    p = tf1_2.add_paragraph()
    p.text = "4. Max FOIR Ratio (≤ 40% | Max 50%):\n   FOIR = (Existing EMIs + Proposed EMI) / Net Income × 100"
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    p = tf1_2.add_paragraph()
    p.text = "AUTOMATED ROUTING ENGINE:\n" \
             "• CIBIL ≥ 750 & FOIR ≤ 40% ➔ [PRE_APPROVED]\n" \
             "• CIBIL < 650 or FOIR > 50% ➔ [HIGH_RISK_EXCEPTION]\n" \
             "• Boundary Criteria Met ➔ [MANUAL_REVIEW]"
    p.font.size = Pt(7.8)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    # ════════════════ COLUMN 2: RBAC & ZERO-TRUST ════════════════
    col2_x = margin_x + col_width + col_gap
    c2_h1 = Inches(4.3)
    c2_h2 = col_height - c2_h1 - Inches(0.18)

    # Card 2.1 Dual-Role Architecture
    tf2_1 = add_card(col2_x, col_top, col_width, c2_h1, "Dual-Role Portal Architecture", "[👥]")
    p = tf2_1.paragraphs[0]
    p.text = "Strict Role-Based Access Control separating customers from underwriters:"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = tf2_1.add_paragraph()
    p.text = "END USER / BORROWER PORTAL (/my-applications):\n" \
             "• Self-Registration & Dual-Tab Login screen.\n" \
             "• Auto-Populated Application Form (/apply) pre-filled from user profile.\n" \
             "• Personal Portfolio View: Borrowers only view applications they submitted.\n" \
             "• Read-Only Audit Timeline: Tracking stages without modification rights.\n" \
             "• Zero Access to Bank Overview or Alerts Queue."
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_BODY

    p = tf2_1.add_paragraph()
    p.text = "BANK OFFICER / ADMIN CONSOLE (/dashboard):\n" \
             "• Global Portfolio Dashboard: Live view of all bank loan applications.\n" \
             "• 4 Live KPI Counters: Total, In-Review, Approved, High-Risk.\n" \
             "• Real-Time Search: Filter by Application ID or Applicant Name.\n" \
             "• Lifecycle Stage Advancement:\n" \
             "  Submitted ➔ KYC Verified ➔ Docs Verified ➔ Underwriting ➔ Approved / Rejected.\n" \
             "• Alerts Queue (/alerts): Immediate review for exception accounts."
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_BODY

    # Card 2.2 Route Protection Matrix
    tf2_2 = add_card(col2_x, col_top + c2_h1 + Inches(0.18), col_width, c2_h2, "Route Security & Tab Isolation", "[🔒]")
    p = tf2_2.paragraphs[0]
    p.text = "ZERO-TRUST ROUTE GUARD MATRIX:\n" \
             "• /login, /register ➔ Public (Zero internal nav tabs shown)\n" \
             "• /dashboard ➔ ADMIN only (Redirects to /login if unauth)\n" \
             "• /my-applications ➔ USER only (Redirects to /login)\n" \
             "• /apply ➔ USER / ADMIN (Redirects to /login)\n" \
             "• /loans/{id} ➔ Owner or ADMIN only\n" \
             "• /alerts ➔ ADMIN only (403 Forbidden for User)"
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_BODY

    p = tf2_2.add_paragraph()
    p.text = "DOM TAB PRIVACY GUARANTEE:\n" \
             "Before logging in, navigation tabs (Dashboard, New Application, Alerts) are completely removed from the DOM to eliminate data leakage."
    p.font.size = Pt(7.8)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED

    # ════════════════ COLUMN 3: JENKINS & SELENIUM ════════════════
    col3_x = margin_x + (col_width + col_gap) * 2
    c3_h1 = Inches(3.5)
    c3_h2 = col_height - c3_h1 - Inches(0.18)

    # Card 3.1 Jenkins Declarative Pipeline
    tf3_1 = add_card(col3_x, col_top, col_width, c3_h1, "Jenkins Declarative Pipeline", "[⚙]")
    p = tf3_1.paragraphs[0]
    p.text = "Automated 6-stage CI/CD execution triggered via Git Webhook on Port 8080:"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = tf3_1.add_paragraph()
    p.text = "STAGE 1: Checkout SCM (Git Feature/Main Branch)\n" \
             "STAGE 2: Unit & Integration Tests (18 Tests - Model/Repo/Service)\n" \
             "STAGE 3: Selenium E2E Quality Gate (9 Chrome Headless Journeys)\n" \
             "STAGE 4: Multi-Stage Docker Build (Lightweight container)\n" \
             "STAGE 5: Continuous Deployment (Port 8081 isolated runtime)\n" \
             "STAGE 6: Actuator Health Check (HTTP 200 OK /actuator/health)"
    p.font.size = Pt(8)
    p.font.color.rgb = TEXT_BODY

    p = tf3_1.add_paragraph()
    p.text = "STRICT QUALITY GATE POLICY:\n" \
             "Any test or Selenium journey failure aborts the pipeline immediately, captures a failure screenshot, and prevents artifact deployment."
    p.font.size = Pt(7.8)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    # Card 3.2 Selenium 9 Automated Journeys
    tf3_2 = add_card(col3_x, col_top + c3_h1 + Inches(0.18), col_width, c3_h2, "Headless Selenium Quality Gate", "[🔬]")
    p = tf3_2.paragraphs[0]
    p.text = "Comprehensive browser automation executing across 9 real user journeys:"
    p.font.size = Pt(8.2)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    journeys = [
        ("J1 (Intake)", "Submits loan form ➔ verifies generated LOAN-XXXX on table."),
        ("J2 (Search)", "Filters live table by applicant name with instant reset."),
        ("J3 (Drill-Down)", "Advances stage to DOCUMENT_VERIFIED ➔ verifies audit log."),
        ("J4 (Rejection)", "Rejects loan with reason code ➔ verifies alert queue entry."),
        ("J5 (Low Risk)", "CIBIL=820, FOIR ~13% ➔ verifies PRE_APPROVED."),
        ("J6 (High Risk)", "CIBIL=580 ➔ verifies auto-tag HIGH_RISK_EXCEPTION."),
        ("J7 (Borrower)", "User login ➔ tracks application ➔ verifies read-only stages."),
        ("J8 (Admin)", "Admin login ➔ approves loan with remarks ➔ reviews alerts."),
        ("J9 (Security)", "Asserts unauthenticated visitor is blocked with zero tabs shown.")
    ]
    for j_tag, j_desc in journeys:
        p = tf3_2.add_paragraph()
        p.text = f"• {j_tag}: {j_desc}"
        p.font.size = Pt(7.3)
        p.font.color.rgb = TEXT_BODY

    p = tf3_2.add_paragraph()
    p.text = "TOTAL SUITE: 27 / 27 Tests Passed (100% Quality Gate Pass Rate)"
    p.font.size = Pt(7.8)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    # ════════════════ COLUMN 4: INFRASTRUCTURE & METRICS ════════════════
    col4_x = margin_x + (col_width + col_gap) * 3
    c4_h1 = Inches(2.2)
    c4_h2 = Inches(2.6)
    c4_h3 = col_height - c4_h1 - c4_h2 - Inches(0.36)

    # Card 4.1 Telemetry & KPI Metrics
    tf4_1 = add_card(col4_x, col_top, col_width, c4_h1, "Telemetry & Portfolio Metrics", "[📊]")
    p = tf4_1.paragraphs[0]
    p.text = "REAL-TIME CI/CD & APPLICATION TELEMETRY:"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = tf4_1.add_paragraph()
    p.text = "• 27 / 27 Automated Tests Passing (Unit, Service, E2E)\n" \
             "• 100% Zero-Defect Quality Gate Pass Rate\n" \
             "• < 15 Seconds Automated Rollback MTTR\n" \
             "• 0 Production Deployment Failures"
    p.font.size = Pt(8)
    p.font.color.rgb = TEXT_BODY

    p = tf4_1.add_paragraph()
    p.text = "• Actuator Endpoint: /actuator/health returns {\"status\":\"UP\"}"
    p.font.size = Pt(7.8)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    # Card 4.2 Docker Multi-Stage & Port Isolation
    tf4_2 = add_card(col4_x, col_top + c4_h1 + Inches(0.18), col_width, c4_h2, "Docker Multi-Stage & Isolation", "[🐳]")
    p = tf4_2.paragraphs[0]
    p.text = "HARDENED MULTI-STAGE CONTAINER ARCHITECTURE:"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = tf4_2.add_paragraph()
    p.text = "• Stage 1 (Builder): Maven 3.9 + Eclipse-Temurin JDK 17.\n" \
             "• Stage 2 (Runtime): Minimal Eclipse-Temurin 17 JRE Alpine (<200MB).\n" \
             "• Security Hardening: Runs under non-root appuser:appgroup."
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_BODY

    p = tf4_2.add_paragraph()
    p.text = "PORT ISOLATION MATRIX:\n" \
             "• Port 8080 ➔ Jenkins CI/CD Controller\n" \
             "• Port 8081 ➔ Loan Tracking Application Service\n" \
             "• Port 22   ➔ SSH / Ansible Automation Node"
    p.font.size = Pt(7.8)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    # Card 4.3 Ansible Disaster Recovery
    tf4_3 = add_card(col4_x, col_top + c4_h1 + c4_h2 + Inches(0.36), col_width, c4_h3, "Ansible IaC & Disaster Recovery", "[🔧]")
    p = tf4_3.paragraphs[0]
    p.text = "IDEMPOTENT PROVISIONING & INSTANT ROLLBACK:"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    p = tf4_3.add_paragraph()
    p.text = "• Ansible Playbook configures OS packages, directories, and dependencies with zero configuration drift (changed=0 on rerun)."
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_BODY

    p = tf4_3.add_paragraph()
    p.text = "• Sub-15s Rollback Automation (Task 14 Deliverable):\n" \
             "  powershell -File scripts\\verify_reliability.ps1\n" \
             "  Restores previous container image loan-portal:v1.0.0 upon failure."
    p.font.size = Pt(7.8)
    p.font.color.rgb = TEXT_BODY

    p = tf4_3.add_paragraph()
    p.text = "DEFAULT DEMO ACCOUNTS:\n" \
             "• Borrower: user / user123 (My Applications Portal)\n" \
             "• Officer:  admin / admin123 (Bank Management Console)"
    p.font.size = Pt(7.8)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    # ── Footer Bar ──
    footer_y = prs.slide_height - footer_height - Inches(0.15)
    footer_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, margin_x, footer_y, content_width, footer_height)
    footer_box.fill.solid()
    footer_box.fill.fore_color.rgb = BG_CARD
    footer_box.line.color.rgb = RGBColor(0, 168, 150)
    footer_box.line.width = Pt(2)

    tf_f = footer_box.text_frame
    tf_f.word_wrap = False
    tf_f.margin_left = Inches(0.2)
    tf_f.margin_top = Inches(0.08)
    tf_f.margin_right = Inches(0.2)
    tf_f.margin_bottom = Inches(0.08)

    pf = tf_f.paragraphs[0]
    pf.text = "Project: Secure CI/CD Pipeline for a Loan Application Tracking Portal  |  Candidate: Vicky Vilas Pukale (Roll: 23102A0033)  |  Repo: github.com/vickypukale/loan-application-tracking-portal"
    pf.font.bold = True
    pf.font.size = Pt(9.5)
    pf.font.color.rgb = TEXT_DARK

    # Save PowerPoint file
    out_path = os.path.abspath("DevOps_Project_A3_Poster.pptx")
    prs.save(out_path)
    print(f"Successfully generated PowerPoint A3 Poster at: {out_path}")
    print(f"File size: {os.path.getsize(out_path)} bytes")

if __name__ == "__main__":
    create_a3_poster()
